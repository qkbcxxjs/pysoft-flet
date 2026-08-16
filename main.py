from __future__ import annotations

import os
import re
import secrets
import shutil
import sqlite3
import sys
from calendar import monthrange
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

import flet as ft
import fitz
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pypinyin import Style, lazy_pinyin


class Services:
    DATA_FOLDER = "qkbc_merchant_data"
    IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
    PREVIEW_SUFFIXES = IMAGE_SUFFIXES | {".pdf"}
    TEMP_SESSION: Path | None = None

    @staticmethod
    def get_app_dir() -> Path:
        if getattr(sys, "frozen", False):
            return Path(sys.executable).resolve().parent
        return Path(__file__).resolve().parent

    @staticmethod
    def get_bundle_dir() -> Path:
        if getattr(sys, "frozen", False):
            return Path(getattr(sys, "_MEIPASS"))
        return Path(__file__).resolve().parent

    @staticmethod
    def get_asset_path(filename: str) -> Path:
        return Services.get_bundle_dir() / "assets" / filename

    @staticmethod
    def get_data_dir() -> Path:
        path = Services.get_app_dir() / Services.DATA_FOLDER
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def get_temp_dir() -> Path:
        path = Services.get_data_dir() / "temp"
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def set_cleanup_temp() -> None:
        cutoff = datetime.now() - timedelta(hours=1)
        for item in Services.get_temp_dir().iterdir():
            if not item.is_dir():
                continue
            try:
                modified = datetime.fromtimestamp(item.stat().st_mtime)
                if modified < cutoff:
                    shutil.rmtree(item)
            except OSError:
                continue

    @staticmethod
    def get_temp_session_dir() -> Path:
        if Services.TEMP_SESSION is None:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            path = Services.get_temp_dir() / stamp
            counter = 1
            while path.exists():
                path = Services.get_temp_dir() / f"{stamp}_{counter}"
                counter += 1
            path.mkdir(parents=True)
            Services.TEMP_SESSION = path
        return Services.TEMP_SESSION

    @staticmethod
    def get_data_path(*parts: str) -> Path:
        return Services.get_data_dir().joinpath(*parts)

    @staticmethod
    def get_relative_path(path: Path) -> str:
        return path.resolve().relative_to(Services.get_data_dir().resolve()).as_posix()

    @staticmethod
    def get_absolute_data_path(relative_path: str) -> Path:
        root = Services.get_data_dir().resolve()
        path = (root / relative_path).resolve()
        if path != root and root not in path.parents:
            raise ValueError("非法文件路径")
        return path

    @staticmethod
    def get_is_image(path: str | Path) -> bool:
        return Path(path).suffix.lower() in Services.IMAGE_SUFFIXES

    @staticmethod
    def get_can_preview(path: str | Path) -> bool:
        return Path(path).suffix.lower() in Services.PREVIEW_SUFFIXES

    @staticmethod
    def get_safe_name(name: str) -> str:
        cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name.strip())
        return cleaned.rstrip(". ") or "附件"

    @staticmethod
    def get_temp_copy(source: str | Path) -> Path:
        source_path = Path(source)
        if not source_path.is_file():
            raise FileNotFoundError("选择的文件不存在")
        target = Services.get_temp_session_dir() / f"{secrets.token_hex(8).upper()}{source_path.suffix.lower()}"
        shutil.copy2(source_path, target)
        return target

    @staticmethod
    def set_image_to_webp(source: str | Path, target: str | Path) -> Path:
        target_path = Path(target)
        target_path.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(source) as image:
            image.seek(0)
            if image.mode not in ("RGB", "RGBA"):
                image = image.convert("RGBA" if "transparency" in image.info else "RGB")
            image.save(target_path, "WEBP", quality=90, method=6)
        return target_path

    @staticmethod
    def set_attachment(source: str | Path, merchant_code: str, file_type: int, display_name: str) -> Path:
        source_path = Path(source)
        suffix = ".webp" if Services.get_is_image(source_path) else source_path.suffix.lower()
        filename = f"{Services.get_safe_name(display_name)}_{secrets.token_hex(3)[:5].upper()}{suffix}"
        target = Services.get_data_path(merchant_code, "attachments", str(file_type), filename)
        target.parent.mkdir(parents=True, exist_ok=True)
        if Services.get_is_image(source_path):
            Services.set_image_to_webp(source_path, target)
        else:
            shutil.copy2(source_path, target)
        return target

    @staticmethod
    def set_transaction_file(
        source: str | Path, merchant_code: str, transaction_id: int,
        file_type: int, display_name: str, type_name: str,
    ) -> Path:
        source_path = Path(source)
        suffix = ".webp" if Services.get_is_image(source_path) else source_path.suffix.lower()
        prefix = display_name if file_type == 10 else type_name
        filename = f"{Services.get_safe_name(prefix)}_{secrets.token_hex(3)[:5].upper()}{suffix}"
        target = Services.get_data_path(merchant_code, "transaction", str(transaction_id), filename)
        target.parent.mkdir(parents=True, exist_ok=True)
        if Services.get_is_image(source_path):
            Services.set_image_to_webp(source_path, target)
        else:
            shutil.copy2(source_path, target)
        return target

    @staticmethod
    def set_remove_file(path: str | Path) -> None:
        file_path = Path(path)
        if file_path.is_file():
            file_path.unlink()

    @staticmethod
    def set_remove_tree(path: str | Path) -> None:
        folder = Path(path)
        if folder.is_dir() and Services.get_data_dir().resolve() in folder.resolve().parents:
            shutil.rmtree(folder)

    @staticmethod
    def set_remove_empty_parents(path: str | Path) -> None:
        root = Services.get_data_dir().resolve()
        current = Path(path).resolve()
        while current != root and root in current.parents:
            try:
                current.rmdir()
            except OSError:
                break
            current = current.parent

    @staticmethod
    def get_pdf_page(path: str | Path, page_index: int) -> tuple[bytes, int]:
        with fitz.open(path) as document:
            total = document.page_count
            page = document.load_page(max(0, min(page_index, total - 1)))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            return pixmap.tobytes("png"), total

    @staticmethod
    def set_open_folder(path: str | Path) -> None:
        target = Path(path)
        folder = target if target.is_dir() else target.parent
        os.startfile(str(folder))

    @staticmethod
    def get_chart_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        names = ["msyhbd.ttc" if bold else "msyh.ttc", "simhei.ttf"]
        for name in names:
            path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / name
            if path.exists():
                return ImageFont.truetype(str(path), size)
        return ImageFont.load_default()

    @staticmethod
    def get_donut_chart(expense: int, income: int, title: str, size: tuple[int, int] = (360, 320)) -> bytes:
        image = Image.new("RGBA", size, "white")
        draw = ImageDraw.Draw(image)
        red, green = "#D64242", "#239B62"
        font, bold = Services.get_chart_font(15), Services.get_chart_font(17, True)
        draw.text((size[0] / 2 - 8, 22), f"支出 ￥{expense / 100:.2f}", fill=red, font=font, anchor="rm")
        draw.text((size[0] / 2 + 8, 22), f"收入 ￥{income / 100:.2f}", fill=green, font=font, anchor="lm")
        box = (75, 58, size[0] - 75, size[1] - 62)
        total = expense + income
        if total:
            split = expense / total * 360
            draw.arc(box, -90, -90 + split, fill=red, width=34)
            draw.arc(box, -90 + split, 270, fill=green, width=34)
        else:
            draw.ellipse(box, outline="#DDE3E6", width=34)
        profit = income - expense
        text = "0" if profit == 0 else f'{"+" if profit > 0 else "-"}￥{abs(profit) / 100:.2f}'
        draw.text((size[0] / 2, 164), text, fill=green if profit > 0 else red if profit < 0 else "#202124",
                  font=bold, anchor="mm")
        draw.text((size[0] / 2, size[1] - 23), title, fill="#202124", font=bold, anchor="mm")
        output = BytesIO(); image.save(output, "PNG")
        return output.getvalue()

    @staticmethod
    def get_trend_chart(points: list[tuple[str, int, int]], title: str, size: tuple[int, int] = (1040, 560)) -> bytes:
        image = Image.new("RGBA", size, "white")
        draw = ImageDraw.Draw(image)
        red, green, orange = "#D64242", "#239B62", "#E78222"
        font, small, bold = Services.get_chart_font(14), Services.get_chart_font(11), Services.get_chart_font(16, True)
        legend_y = 22
        draw.rectangle((size[0] / 2 - 220, legend_y - 7, size[0] / 2 - 204, legend_y + 7), fill=red)
        draw.text((size[0] / 2 - 196, legend_y), "支出", fill=red, font=font, anchor="lm")
        draw.rectangle((size[0] / 2 - 80, legend_y - 7, size[0] / 2 - 64, legend_y + 7), fill=green)
        draw.text((size[0] / 2 - 56, legend_y), "收入", fill=green, font=font, anchor="lm")
        draw.line((size[0] / 2 + 55, legend_y, size[0] / 2 + 83, legend_y), fill=orange, width=3)
        draw.ellipse((size[0] / 2 + 65, legend_y - 5, size[0] / 2 + 75, legend_y + 5), fill="white", outline=orange, width=2)
        draw.text((size[0] / 2 + 95, legend_y), "利润", fill=orange, font=font, anchor="lm")
        left, right, top, bottom = 92, size[0] - 25, 62, size[1] - 112
        values = [value for _, expense, income in points for value in (-expense, income, income - expense)] or [0]
        limit = max(100, max(abs(x) for x in values))
        zero = (top + bottom) / 2
        draw.line((left, top, left, bottom), fill="#89949A", width=2)
        tick_value = max(1, round(limit / 100 / 4))
        for tick in range(-4, 5):
            y = zero - tick * (bottom - top) / 8
            draw.line((left, y, right, y), fill="#D9E0E3" if tick else "#89949A", width=2 if tick == 0 else 1)
            label_value = tick * tick_value
            draw.text((left - 8, y), str(label_value), fill="#59656B", font=small, anchor="rm")
        count = max(1, len(points)); step = (right - left) / count; line_points = []
        for index, (label, expense, income) in enumerate(points):
            x = left + step * (index + 0.5); scale = (bottom - top) / 2 / limit
            expense_y, income_y = zero + expense * scale, zero - income * scale
            width = max(2, min(12, step * 0.24))
            draw.rectangle((x - width - 1, zero, x - 1, expense_y), fill=red)
            draw.rectangle((x + 1, income_y, x + width + 1, zero), fill=green)
            profit_y = zero - (income - expense) * scale
            line_points.append((x, profit_y))
            label_image = Image.new("RGBA", (90, 18), (255, 255, 255, 0))
            ImageDraw.Draw(label_image).text((45, 9), label, fill="#59656B", font=small, anchor="mm")
            label_image = label_image.rotate(90, expand=True)
            image.alpha_composite(label_image, (int(x - label_image.width / 2), int(bottom + 6)))
        if len(line_points) > 1:
            draw.line(line_points, fill=orange, width=3)
        for x, y in line_points:
            draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill="white", outline=orange, width=2)
        draw.text((size[0] / 2, size[1] - 16), title, fill="#202124", font=bold, anchor="mm")
        output = BytesIO(); image.save(output, "PNG")
        return output.getvalue()


class QKBCASSqlite:
    def __init__(self) -> None:
        self.db_path = Services.get_data_path("qkbcas.db")
        Services.set_cleanup_temp()
        Services.get_temp_session_dir()
        self.set_create_tables()

    def get_connection(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def set_create_tables(self) -> None:
        with self.get_connection() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS merchant_info (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    code TEXT NOT NULL,
                    address TEXT NOT NULL DEFAULT '',
                    legal_name TEXT NOT NULL DEFAULT '',
                    legal_phone TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE TABLE IF NOT EXISTS merchant_file (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    merchant_id INTEGER NOT NULL,
                    name TEXT NOT NULL,
                    type INTEGER NOT NULL CHECK(type BETWEEN 0 AND 3),
                    path TEXT NOT NULL,
                    FOREIGN KEY(merchant_id) REFERENCES merchant_info(id) ON DELETE CASCADE
                );
                CREATE TABLE IF NOT EXISTS payee_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL UNIQUE,
                    count INTEGER NOT NULL DEFAULT 0,
                    last_used TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    pinyin TEXT NOT NULL DEFAULT '',
                    category INTEGER NOT NULL DEFAULT 0 CHECK(category BETWEEN 0 AND 2)
                );
                CREATE TABLE IF NOT EXISTS transaction_info (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    merchant_id INTEGER NOT NULL,
                    type INTEGER NOT NULL CHECK(type IN (0, 1)),
                    amount INTEGER NOT NULL,
                    description TEXT NOT NULL DEFAULT '',
                    way INTEGER NOT NULL CHECK(way BETWEEN 0 AND 6),
                    way_remark TEXT NOT NULL DEFAULT '',
                    payer TEXT NOT NULL,
                    recipient TEXT NOT NULL,
                    payer_id INTEGER,
                    recipient_id INTEGER,
                    date TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(merchant_id) REFERENCES merchant_info(id) ON DELETE CASCADE,
                    FOREIGN KEY(payer_id) REFERENCES payee_history(id) ON DELETE SET NULL,
                    FOREIGN KEY(recipient_id) REFERENCES payee_history(id) ON DELETE SET NULL
                );
                CREATE TABLE IF NOT EXISTS transaction_files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    transaction_id INTEGER NOT NULL,
                    original_name TEXT NOT NULL,
                    name TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    type INTEGER NOT NULL CHECK(type BETWEEN 0 AND 10),
                    path TEXT NOT NULL,
                    FOREIGN KEY(transaction_id) REFERENCES transaction_info(id) ON DELETE CASCADE
                );
                """
            )
            rows = connection.execute("SELECT id,name,pinyin FROM payee_history").fetchall()
            for row in rows:
                pinyin = "".join(lazy_pinyin(row["name"], style=Style.NORMAL)).casefold()
                if row["pinyin"] != pinyin:
                    connection.execute("UPDATE payee_history SET pinyin=? WHERE id=?", (pinyin, row["id"]))

    def get_merchants(self) -> list[dict[str, Any]]:
        with self.get_connection() as connection:
            rows = connection.execute("SELECT * FROM merchant_info ORDER BY id DESC").fetchall()
        return [dict(row) for row in rows]

    def get_merchant(self, merchant_id: int) -> dict[str, Any] | None:
        with self.get_connection() as connection:
            row = connection.execute("SELECT * FROM merchant_info WHERE id = ?", (merchant_id,)).fetchone()
        return dict(row) if row else None

    def get_files(self, merchant_id: int) -> list[dict[str, Any]]:
        with self.get_connection() as connection:
            rows = connection.execute(
                "SELECT * FROM merchant_file WHERE merchant_id = ? ORDER BY id", (merchant_id,)
            ).fetchall()
        return [dict(row) for row in rows]

    def set_save_merchant(self, merchant_id: int | None, values: dict[str, str]) -> int:
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self.get_connection() as connection:
            if merchant_id is None:
                cursor = connection.execute(
                    """INSERT INTO merchant_info
                    (name, code, address, legal_name, legal_phone, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?)""",
                    (*values.values(), now, now),
                )
                return int(cursor.lastrowid)
            connection.execute(
                """UPDATE merchant_info SET name=?, code=?, address=?, legal_name=?,
                legal_phone=?, updated_at=? WHERE id=?""",
                (*values.values(), now, merchant_id),
            )
            return merchant_id

    def set_replace_files(self, merchant_id: int, files: list[dict[str, Any]]) -> None:
        with self.get_connection() as connection:
            connection.execute("DELETE FROM merchant_file WHERE merchant_id = ?", (merchant_id,))
            connection.executemany(
                "INSERT INTO merchant_file (merchant_id, name, type, path) VALUES (?, ?, ?, ?)",
                [(merchant_id, item["name"], item["type"], item["path"]) for item in files],
            )

    def set_delete_merchant(self, merchant_id: int) -> None:
        with self.get_connection() as connection:
            connection.execute("DELETE FROM merchant_info WHERE id = ?", (merchant_id,))

    def get_transactions(self, merchant_id: int) -> list[dict[str, Any]]:
        with self.get_connection() as connection:
            rows = connection.execute(
                "SELECT * FROM transaction_info WHERE merchant_id=? ORDER BY date DESC, id DESC", (merchant_id,)
            ).fetchall()
        return [dict(row) for row in rows]

    def get_transaction(self, transaction_id: int) -> dict[str, Any] | None:
        with self.get_connection() as connection:
            row = connection.execute("SELECT * FROM transaction_info WHERE id=?", (transaction_id,)).fetchone()
        return dict(row) if row else None

    def get_transaction_files(self, transaction_id: int) -> list[dict[str, Any]]:
        with self.get_connection() as connection:
            rows = connection.execute(
                "SELECT * FROM transaction_files WHERE transaction_id=? ORDER BY id", (transaction_id,)
            ).fetchall()
        return [dict(row) for row in rows]

    def set_payee_history(self, name: str, category: int) -> int:
        normalized = " ".join(name.strip().split())
        pinyin = "".join(lazy_pinyin(normalized, style=Style.NORMAL)).casefold()
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self.get_connection() as connection:
            row = connection.execute("SELECT id, category FROM payee_history WHERE name=?", (normalized,)).fetchone()
            if row:
                merged_category = row["category"] if row["category"] == category else 0
                connection.execute(
                    "UPDATE payee_history SET count=count+1,last_used=?,pinyin=?,category=? WHERE id=?",
                    (now, pinyin, merged_category, row["id"]),
                )
                return int(row["id"])
            cursor = connection.execute(
                "INSERT INTO payee_history(name,count,last_used,pinyin,category) VALUES(?,1,?,?,?)",
                (normalized, now, pinyin, category),
            )
            return int(cursor.lastrowid)

    def get_payee_suggestions(self, query: str, category: int, limit: int = 8) -> list[dict[str, Any]]:
        keyword = query.strip().casefold()
        now = datetime.now()
        with self.get_connection() as connection:
            rows = connection.execute(
                "SELECT * FROM payee_history WHERE category IN (0,?) AND (name LIKE ? OR pinyin LIKE ?)",
                (category, f"%{keyword}%", f"%{keyword}%"),
            ).fetchall()
        scored = []
        for row in rows:
            item = dict(row)
            text, index = item["name"].casefold(), item["pinyin"].casefold()
            match = 3 if text == keyword else 2 if text.startswith(keyword) or index.startswith(keyword) else 1
            try:
                days = max(0, (now - datetime.fromisoformat(item["last_used"])).days)
            except ValueError:
                days = 365
            recent = max(0.0, 1 - days / 30)
            item["score"] = match * 100 + recent * 20 + item["count"]
            scored.append(item)
        return sorted(scored, key=lambda x: x["score"], reverse=True)[:limit]

    def set_save_transaction(self, transaction_id: int | None, values: dict[str, Any]) -> int:
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self.get_connection() as connection:
            params = (values["merchant_id"], values["type"], values["amount"], values["description"],
                      values["way"], values["way_remark"], values["payer"], values["recipient"],
                      values["payer_id"], values["recipient_id"], values["date"])
            if transaction_id is None:
                cursor = connection.execute(
                    """INSERT INTO transaction_info
                    (merchant_id,type,amount,description,way,way_remark,payer,recipient,payer_id,recipient_id,date,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""", (*params, now, now))
                return int(cursor.lastrowid)
            connection.execute(
                """UPDATE transaction_info SET merchant_id=?,type=?,amount=?,description=?,way=?,way_remark=?,
                payer=?,recipient=?,payer_id=?,recipient_id=?,date=?,updated_at=? WHERE id=?""",
                (*params, now, transaction_id),
            )
            return transaction_id

    def set_replace_transaction_files(self, transaction_id: int, files: list[dict[str, Any]]) -> None:
        with self.get_connection() as connection:
            connection.execute("DELETE FROM transaction_files WHERE transaction_id=?", (transaction_id,))
            connection.executemany(
                """INSERT INTO transaction_files
                (transaction_id,original_name,name,size,type,path) VALUES(?,?,?,?,?,?)""",
                [(transaction_id, x["original_name"], x["name"], x["size"], x["type"], x["path"]) for x in files],
            )

    def get_filtered_transactions(
        self, merchant_id: int, start: str | None, end: str | None,
        transaction_type: int, way: int, order: str,
    ) -> list[dict[str, Any]]:
        conditions, params = ["merchant_id=?"], [merchant_id]
        if start:
            conditions.append("date>=?")
            params.append(start)
        if end:
            conditions.append("date<=?")
            params.append(end)
        if transaction_type in (0, 1):
            conditions.append("type=?")
            params.append(transaction_type)
        if way in range(7):
            conditions.append("way=?")
            params.append(way)
        order_sql = {
            "date_desc": "date DESC, amount DESC",
            "date_asc": "date ASC, amount DESC",
            "amount_desc": "amount DESC, date DESC",
            "amount_asc": "amount ASC, date DESC",
        }[order]
        sql = f"SELECT * FROM transaction_info WHERE {' AND '.join(conditions)} ORDER BY {order_sql}, id DESC"
        with self.get_connection() as connection:
            rows = connection.execute(sql, params).fetchall()
        return [dict(row) for row in rows]

    def set_delete_transaction(self, transaction_id: int) -> None:
        with self.get_connection() as connection:
            connection.execute("DELETE FROM transaction_info WHERE id=?", (transaction_id,))


class Store:
    def __init__(self) -> None:
        self.selected_nav_index = 0
        self.selected_merchant_id: int | None = None

    def get_selected_nav_index(self) -> int:
        return self.selected_nav_index

    def set_selected_nav_index(self, index: int) -> None:
        self.selected_nav_index = index


class QKBCASApp:
    WINDOW_WIDTH, WINDOW_HEIGHT = 1320, 850
    WINDOW_MIN_WIDTH, WINDOW_MIN_HEIGHT = 1050, 700
    NAV_ITEMS = (
        ("商户管理", "icon-store.svg"),
        ("所有账目", "icon-accounts.svg"),
        ("收支记账", "icon-bookkeeping.svg"),
        ("账目统计", "icon-workbench.svg"),
    )
    FILE_TYPES = {0: "营业执照", 1: "许可证", 2: "健康证", 3: "其他"}
    TRANSACTION_FILE_TYPES = {
        0: "发票", 1: "付款截图", 2: "订单截图", 3: "收据", 4: "银行凭证", 5: "车票机票",
        6: "完税凭证", 7: "工资表", 8: "出差申请单", 9: "考勤记录", 10: "其他",
    }
    TRANSACTION_WAYS = {
        0: "现金", 1: "支付宝", 2: "微信", 3: "对公转账", 4: "现金+支付宝",
        5: "现金+微信", 6: "其他",
    }

    def __init__(self, page: ft.Page) -> None:
        self.page, self.store, self.db = page, Store(), QKBCASSqlite()
        self.nav_panel, self.content_panel = ft.Container(), ft.Container(expand=True)
        self.nav_controls: list[ft.Container] = []
        self.file_cards: list[dict[str, Any]] = []
        self.transaction_file_cards: list[dict[str, Any]] = []
        self.selected_transaction_id: int | None = None
        self.name_field = ft.TextField(label="商户名称", hint_text="与营业执照保持一致", expand=True)
        self.code_field = ft.TextField(label="统一社会信用代码", hint_text="与营业执照保持一致", expand=True)
        self.address_field = ft.TextField(label="地址", hint_text="与营业执照保持一致", expand=True)
        self.legal_name_field = ft.TextField(label="法人姓名", hint_text="请输入法人姓名", expand=True)
        self.legal_phone_field = ft.TextField(label="法人手机号", hint_text="请输入法人手机号", expand=True)
        self.merchant_dropdown = ft.Dropdown(label="选择商户", expand=True, on_select=self.set_select_merchant)
        self.attachments_row = ft.ResponsiveRow(spacing=18, run_spacing=18)
        self.transaction_dropdown = ft.Dropdown(label="账目记录", expand=True, on_select=self.set_select_transaction)
        self.transaction_merchant_dropdown = ft.Dropdown(label="记账商户", expand=True, on_select=self.set_transaction_merchant)
        self.amount_field = ft.TextField(label="金额（元）", hint_text="0.00", expand=True)
        self.transaction_type = ft.RadioGroup(value="0", content=ft.Row(controls=[
            ft.Radio(value="0", label="支出", label_style=ft.TextStyle(color="#C62828")),
            ft.Radio(value="1", label="收入", label_style=ft.TextStyle(color="#18864B")),
        ]))
        self.payer_field = ft.TextField(label="付款方", label_style=ft.TextStyle(color="#C62828"), expand=True,
                                        on_change=lambda e: self.set_update_payee_popup(e, 1))
        self.recipient_field = ft.TextField(label="收款方", label_style=ft.TextStyle(color="#18864B"), expand=True,
                                            on_change=lambda e: self.set_update_payee_popup(e, 2))
        self.payer_popup = ft.PopupMenuButton(icon=ft.Icons.ARROW_DROP_DOWN, tooltip="付款方历史")
        self.recipient_popup = ft.PopupMenuButton(icon=ft.Icons.ARROW_DROP_DOWN, tooltip="收款方历史")
        self.way_dropdown = ft.Dropdown(label="交易方式", value="0", expand=True,
            options=[ft.DropdownOption(key=str(k), text=v) for k, v in self.TRANSACTION_WAYS.items()],
            on_select=self.set_way_changed)
        self.way_remark_field = ft.TextField(label="交易方式备注", disabled=True, expand=True)
        self.date_field = ft.TextField(label="交易日期", value=date.today().isoformat(), hint_text="YYYY-MM-DD", expand=True)
        self.description_field = ft.TextField(label="账单描述", multiline=True, min_lines=2, max_lines=5)
        self.transaction_files_row = ft.ResponsiveRow(spacing=18, run_spacing=18)
        self.accounts_rows: list[dict[str, Any]] = []
        self.accounts_page = 1
        self.accounts_merchant = ft.Dropdown(label="选择商户", expand=True, on_select=self.set_accounts_merchant)
        self.accounts_start = ft.TextField(label="开始日期", hint_text="YYYY-MM-DD", expand=True)
        self.accounts_end = ft.TextField(label="结束日期", hint_text="YYYY-MM-DD", expand=True)
        self.accounts_sort = ft.Dropdown(label="排序方式", value="date_desc", expand=True, options=[
            ft.DropdownOption(key="date_desc", text="日期（降序）"), ft.DropdownOption(key="date_asc", text="日期（升序）"),
            ft.DropdownOption(key="amount_desc", text="金额（降序）"), ft.DropdownOption(key="amount_asc", text="金额（升序）")])
        self.accounts_type = ft.Dropdown(label="交易类型", value="2", expand=True, options=[
            ft.DropdownOption(key="0", text="支出"), ft.DropdownOption(key="1", text="收入"), ft.DropdownOption(key="2", text="收支")])
        self.accounts_way = ft.Dropdown(label="交易方式", value="7", expand=True,
            options=[ft.DropdownOption(key=str(k), text=v) for k, v in self.TRANSACTION_WAYS.items()] +
                    [ft.DropdownOption(key="7", text="全部")])
        self.accounts_stats = ft.Row(spacing=16)
        self.accounts_table = ft.Column(scroll=ft.ScrollMode.AUTO)
        self.accounts_pager = ft.Row(alignment=ft.MainAxisAlignment.CENTER)
        current_year = date.today().year
        self.stats_merchant = ft.Dropdown(label="选择商户", expand=True, on_select=self.set_stats_merchant)
        self.stats_start = ft.TextField(label="开始日期", hint_text="YYYY-MM-DD", width=180)
        self.stats_end = ft.TextField(label="结束日期", hint_text="YYYY-MM-DD", width=180)
        self.stats_year = ft.Dropdown(label="年份", value=str(current_year), width=130,
            options=[ft.DropdownOption(key=str(year), text=str(year)) for year in range(current_year, current_year - 15, -1)])
        self.stats_quarter = ft.Dropdown(label="季度", width=216, on_select=self.set_query_stats)
        self.stats_cards = ft.Row(spacing=16)
        self.stats_donuts = ft.ResponsiveRow(spacing=12, run_spacing=12)
        self.stats_trends = ft.Column(spacing=22)
        self.stats_data: dict[str, Any] = {}
        self.pdf_preview_box: ft.Container | None = None
        self.pdf_preview_image: ft.Image | None = None
        self.pdf_preview_viewer: ft.InteractiveViewer | None = None

    async def get_setup_window(self) -> None:
        self.page.title = "青科博财税管理软件 V1.3"
        self.page.padding = self.page.spacing = 0
        self.page.bgcolor, self.page.theme_mode = "#F4F6F8", ft.ThemeMode.LIGHT
        self.page.window.width, self.page.window.height = self.WINDOW_WIDTH, self.WINDOW_HEIGHT
        self.page.window.min_width, self.page.window.min_height = self.WINDOW_MIN_WIDTH, self.WINDOW_MIN_HEIGHT
        self.page.window.icon = str(Services.get_asset_path("soft_logo.ico").resolve())
        self.page.theme = ft.Theme(color_scheme_seed="#FD5E0F", font_family="Microsoft YaHei")
        await self.page.window.center()

    def get_nav_width(self) -> int:
        return int((self.page.width or self.page.window.width or self.WINDOW_WIDTH) * 0.2)

    def get_nav_item(self, index: int, label: str, icon: str) -> ft.Container:
        selected = index == self.store.get_selected_nav_index()
        return ft.Container(
            height=58, margin=ft.Margin(18, 4, 18, 4), padding=ft.Padding(18, 0, 18, 0),
            border_radius=6, bgcolor="#E8F2ED" if selected else ft.Colors.TRANSPARENT, ink=True,
            on_click=lambda _e, i=index: self.set_active_view(i),
            content=ft.Row(spacing=15, controls=[
                ft.Image(src=icon, width=24, height=24, fit=ft.BoxFit.CONTAIN),
                ft.Text(label, size=16, weight=ft.FontWeight.W_600 if selected else ft.FontWeight.W_400,
                        color="#176B45" if selected else "#36434A"),
            ]),
        )

    def get_nav_panel(self) -> ft.Container:
        self.nav_controls = [self.get_nav_item(i, *item) for i, item in enumerate(self.NAV_ITEMS)]
        return ft.Container(
            width=self.get_nav_width(), bgcolor="#FFFFFF", border=ft.Border(right=ft.BorderSide(1, "#E5E9EC")),
            padding=ft.Padding(0, 34, 0, 24), content=ft.Column(spacing=22, controls=[
                ft.Container(height=90, padding=ft.Padding(28, 0, 28, 0), alignment=ft.Alignment.CENTER,
                             content=ft.Image(src="logo-nav-top.svg", fit=ft.BoxFit.CONTAIN, aspect_ratio=23 / 10)),
                ft.Column(spacing=0, controls=self.nav_controls),
            ]),
        )

    def get_field_style(self) -> None:
        for field in [self.name_field, self.code_field, self.address_field, self.legal_name_field, self.legal_phone_field]:
            field.border_radius, field.bgcolor = 7, "#FFFFFF"

    def empty_state(
        self, icon=ft.Icons.INBOX, title: str = "暂无数据", description: str = "当前无数据，可前往创建",
        action_label: str | None = None, action_icon=None, action=None,
    ) -> ft.Container:
        wrapped_description = "\n".join(description[index:index + 15] for index in range(0, len(description), 15))
        controls: list[ft.Control] = [
            ft.Icon(icon, size=54, color="#87939A"),
            ft.Text(title, size=16, weight=ft.FontWeight.W_600, text_align=ft.TextAlign.CENTER),
            ft.Text(wrapped_description, size=16, text_align=ft.TextAlign.CENTER, width=240),
        ]
        if action_label is not None:
            controls.append(ft.Button(action_label, icon=action_icon, on_click=action))
        return ft.Container(expand=True, alignment=ft.Alignment.CENTER, padding=40,
            content=ft.Column(tight=True, spacing=10, horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                              controls=controls))

    def get_merchant_view(self) -> ft.Control:
        self.get_field_style()
        self.set_refresh_merchants()
        return ft.Column(expand=True, scroll=ft.ScrollMode.AUTO, spacing=18, controls=[
            ft.Row(controls=[self.merchant_dropdown, ft.Button("新增", icon=ft.Icons.ADD, on_click=self.set_new_merchant)]),
            ft.Divider(),
            ft.Row(spacing=16, controls=[self.name_field, self.code_field]),
            ft.Row(controls=[self.address_field]),
            ft.Row(spacing=16, controls=[self.legal_name_field, self.legal_phone_field]),
            ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.SPACE_BETWEEN, controls=[
                ft.Text("相关附件", size=20, weight=ft.FontWeight.W_600),
                ft.Button("新增", icon=ft.Icons.ADD, on_click=lambda _e: self.set_add_file_card()),
            ]),
            self.attachments_row,
            ft.Divider(),
            ft.Row(spacing=12, alignment=ft.MainAxisAlignment.END, controls=[
                ft.Button("上传", icon=ft.Icons.UPLOAD_FILE, bgcolor="#176B45", color="#FFFFFF", on_click=self.set_save),
                ft.OutlinedButton("删除", icon=ft.Icons.DELETE, on_click=self.set_request_delete),
            ]),
        ])

    def get_placeholder_view(self, index: int) -> ft.Control:
        title, icon = self.NAV_ITEMS[index]
        return ft.Column(expand=True, controls=[
            ft.Row(controls=[ft.Image(src=icon, width=30, height=30), ft.Text(title, size=25, weight=ft.FontWeight.W_600)]),
            ft.Divider(), ft.Container(expand=True, alignment=ft.Alignment.CENTER, content=ft.Text(f"{title}模块", size=18)),
        ])

    def get_transaction_view(self) -> ft.Control:
        self.set_refresh_transaction_merchants()
        self.set_refresh_transactions()
        self.set_update_payee_popup(None, 1)
        self.set_update_payee_popup(None, 2)
        return ft.Column(expand=True, scroll=ft.ScrollMode.AUTO, spacing=18, controls=[
            ft.Row(spacing=12, controls=[
                self.transaction_merchant_dropdown,
                self.transaction_dropdown,
                ft.Button("新增", icon=ft.Icons.ADD, on_click=self.set_new_transaction),
            ]),
            ft.Divider(),
            ft.Row(spacing=16, controls=[
                self.amount_field,
                ft.Container(expand=True, padding=ft.Padding(8, 0, 0, 0), content=self.transaction_type),
            ]),
            ft.Row(spacing=16, controls=[
                ft.Row(expand=True, spacing=0, controls=[self.payer_field, self.payer_popup]),
                ft.Row(expand=True, spacing=0, controls=[self.recipient_field, self.recipient_popup]),
            ]),
            ft.Row(spacing=16, controls=[self.way_dropdown, self.way_remark_field, self.date_field]),
            self.description_field,
            ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.SPACE_BETWEEN, controls=[
                ft.Text("相关凭证", size=20, weight=ft.FontWeight.W_600),
                ft.Button("新增", icon=ft.Icons.ADD, on_click=lambda _e: self.set_add_transaction_file_card()),
            ]),
            self.transaction_files_row,
            ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.END, controls=[
                ft.Button("保存", icon=ft.Icons.SAVE, bgcolor="#176B45", color="#FFFFFF",
                          on_click=self.set_save_transaction),
            ]),
        ])

    @staticmethod
    def get_shift_months(value: date, months: int) -> date:
        month_index = value.year * 12 + value.month - 1 + months
        year, month = divmod(month_index, 12)
        month += 1
        return date(year, month, min(value.day, monthrange(year, month)[1]))

    def get_valid_filter_date(self, value: str | None) -> date | None:
        try:
            parsed = date.fromisoformat((value or "").strip())
            return parsed if parsed <= date.today() else None
        except ValueError:
            return None

    def get_date_menu(self, target: str) -> ft.PopupMenuButton:
        choices = [("当日", 0, False), ("一日", 1, False), ("三日", 3, False), ("七日", 7, False),
                   ("一月", 1, True), ("三月", 3, True), ("六月", 6, True), ("十二月", 12, True)]
        items = []
        for label, amount, is_month in choices:
            prefix = "前" if target == "start" and amount else "后" if target == "end" and amount else ""

            def get_click(delta: int, month_mode: bool):
                def set_click(_event: ft.ControlEvent) -> None:
                    start = self.get_valid_filter_date(self.accounts_start.value)
                    end = self.get_valid_filter_date(self.accounts_end.value)
                    base = end or date.today() if target == "start" else start or date.today()
                    signed = -delta if target == "start" else delta
                    result = self.get_shift_months(base, signed) if month_mode else base + timedelta(days=signed)
                    if result <= date.today():
                        field = self.accounts_start if target == "start" else self.accounts_end
                        field.value = result.isoformat()
                        self.page.update()
                return set_click

            disabled = False
            if target == "end" and amount:
                start = self.get_valid_filter_date(self.accounts_start.value)
                if not start:
                    disabled = True
                else:
                    candidate = self.get_shift_months(start, amount) if is_month else start + timedelta(days=amount)
                    disabled = candidate > date.today()
            items.append(ft.PopupMenuItem(content=f"{prefix}{label}", disabled=disabled,
                                          on_click=get_click(amount, is_month)))
        return ft.PopupMenuButton(icon=ft.Icons.ARROW_DROP_DOWN,
                                  tooltip="推导开始日期" if target == "start" else "推导结束日期", items=items)

    def get_accounts_view(self) -> ft.Control:
        self.set_refresh_accounts_merchants()
        self.set_query_accounts(update=False)
        return ft.Column(expand=True, scroll=ft.ScrollMode.AUTO, spacing=16, controls=[
            ft.Row(controls=[self.accounts_merchant,
                ft.Button("新增", icon=ft.Icons.ADD, on_click=self.set_jump_new_merchant)]),
            ft.Divider(),
            ft.Row(spacing=8, controls=[
                ft.Row(expand=True, spacing=0, controls=[self.accounts_start, self.get_date_menu("start")]),
                ft.Row(expand=True, spacing=0, controls=[self.accounts_end, self.get_date_menu("end")]),
                self.accounts_sort, self.accounts_type, self.accounts_way,
                ft.Button("查询", icon=ft.Icons.SEARCH, on_click=self.set_query_accounts),
            ]),
            ft.Divider(), self.accounts_stats, ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.SPACE_BETWEEN, controls=[
                ft.Text("账目一览表", size=20, weight=ft.FontWeight.W_600),
                ft.Button("新增", icon=ft.Icons.ADD, on_click=self.set_jump_new_transaction),
            ]),
            self.accounts_table, self.accounts_pager, ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.END, controls=[
                ft.Button("导出 Excel", icon=ft.Icons.DOWNLOAD, on_click=self.set_export_accounts),
            ]),
        ])

    def set_refresh_accounts_merchants(self) -> None:
        self.accounts_merchant.options = [ft.DropdownOption(key=str(x["id"]), text=f'{x["name"]}<{x["code"]}>')
            for x in self.db.get_merchants()]
        self.accounts_merchant.value = str(self.store.selected_merchant_id) if self.store.selected_merchant_id else None

    def set_accounts_merchant(self, event: ft.ControlEvent) -> None:
        self.store.selected_merchant_id = int(event.control.value) if event.control.value else None
        self.accounts_page = 1
        self.set_query_accounts()

    def set_query_accounts(self, _event: ft.ControlEvent | None = None, update: bool = True) -> None:
        start, end = self.get_valid_filter_date(self.accounts_start.value), self.get_valid_filter_date(self.accounts_end.value)
        if not start:
            self.accounts_start.value = ""
        if not end:
            self.accounts_end.value = ""
        if start and end and start > end:
            start, end = end, start
            self.accounts_start.value, self.accounts_end.value = start.isoformat(), end.isoformat()
        self.stats_start.value, self.stats_end.value = self.accounts_start.value, self.accounts_end.value
        self.accounts_rows = self.db.get_filtered_transactions(self.store.selected_merchant_id,
            start.isoformat() if start else None, end.isoformat() if end else None,
            int(self.accounts_type.value), int(self.accounts_way.value), self.accounts_sort.value
        ) if self.store.selected_merchant_id else []
        self.accounts_page = 1
        self.set_render_accounts()
        if update and self.page.controls:
            self.page.update()

    def get_stat_card(self, title: str, value: str, color: str) -> ft.Container:
        return ft.Container(expand=True, height=100, bgcolor="#FFFFFF", border=ft.Border.all(1, "#E1E5E8"),
            border_radius=7, alignment=ft.Alignment.CENTER, content=ft.Column(tight=True,
                horizontal_alignment=ft.CrossAxisAlignment.CENTER, controls=[
                    ft.Text(title, weight=ft.FontWeight.W_600, size=24),
                    ft.Text(value, size=22, weight=ft.FontWeight.BOLD, color=color),
                ]))

    def set_render_accounts(self) -> None:
        expense = sum(x["amount"] for x in self.accounts_rows if x["type"] == 0)
        income = sum(x["amount"] for x in self.accounts_rows if x["type"] == 1)
        profit = income - expense
        profit_text = f'{"+" if profit > 0 else "-" if profit < 0 else ""}￥{abs(profit) / 100:.2f}' if profit else "0"
        self.accounts_stats.controls = [
            self.get_stat_card("已支出", f"-￥{expense / 100:.2f}", "#C62828"),
            self.get_stat_card("已收入", f"+￥{income / 100:.2f}", "#18864B"),
            self.get_stat_card("总利润", profit_text, "#18864B" if profit > 0 else "#C62828" if profit < 0 else "#202124"),
        ]
        total_pages = max(1, (len(self.accounts_rows) + 14) // 15)
        self.accounts_page = min(self.accounts_page, total_pages)
        start = (self.accounts_page - 1) * 15
        page_rows = self.accounts_rows[start:start + 15]
        table_rows = []
        for number, item in enumerate(page_rows, start=start + 1):
            description = (item["description"] or "").replace("\r\n", "|").replace("\n", "|").replace("\r", "|")
            description = description if len(description) <= 15 else description[:15] + "..."
            way = item["way_remark"] if item["way"] == 6 else self.TRANSACTION_WAYS[item["way"]]
            party = item["recipient"] if item["type"] == 0 else item["payer"]
            color, sign = ("#C62828", "-") if item["type"] == 0 else ("#18864B", "+")
            table_rows.append(ft.DataRow(cells=[
                ft.DataCell(ft.Text(str(number))), ft.DataCell(ft.Text("支出" if item["type"] == 0 else "收入")),
                ft.DataCell(ft.Text(item["date"])), ft.DataCell(ft.Text(f"{sign}￥{item['amount'] / 100:.2f}", color=color)),
                ft.DataCell(ft.Text(description)), ft.DataCell(ft.Text(way)), ft.DataCell(ft.Text(party)),
                ft.DataCell(ft.IconButton(ft.Icons.MORE_HORIZ, tooltip="预览/修改/删除",
                                          on_click=lambda _e, row=item: self.set_accounts_detail(row))),
            ]))
        if not table_rows:
            self.accounts_table.controls = [self.empty_state(
                description="当前无账目记录，可前往创建", action_label="新增账目",
                action_icon=ft.Icons.ADD, action=self.set_jump_new_transaction)]
        else:
            self.accounts_table.controls = [ft.Row(alignment=ft.MainAxisAlignment.CENTER,
                scroll=ft.ScrollMode.AUTO, controls=[ft.DataTable(
            columns=[ft.DataColumn(ft.Text(x)) for x in
                     ["序号", "交易类型", "交易日期", "金额", "说明", "交易方式", "收/付款方", "操作"]],
            rows=table_rows, column_spacing=24, heading_row_color="#EDF1F3")])]

        def set_page(value: int):
            def change(_event: ft.ControlEvent) -> None:
                self.accounts_page = value
                self.set_render_accounts()
                self.page.update()
            return change

        buttons: list[ft.Control] = []
        if self.accounts_page > 1:
            buttons.extend([ft.IconButton(ft.Icons.FIRST_PAGE, tooltip="开始页", on_click=set_page(1)),
                            ft.IconButton(ft.Icons.NAVIGATE_BEFORE, tooltip="上一页", on_click=set_page(self.accounts_page - 1))])
        buttons.append(ft.Text(f"{self.accounts_page}/{total_pages}"))
        if self.accounts_page < total_pages:
            buttons.extend([ft.IconButton(ft.Icons.NAVIGATE_NEXT, tooltip="下一页", on_click=set_page(self.accounts_page + 1)),
                            ft.IconButton(ft.Icons.LAST_PAGE, tooltip="尾页", on_click=set_page(total_pages))])
        self.accounts_pager.controls = buttons

    def set_accounts_detail(self, item: dict[str, Any]) -> None:
        merchant = self.db.get_merchant(item["merchant_id"])
        files = self.db.get_transaction_files(item["id"])
        way = item["way_remark"] if item["way"] == 6 else self.TRANSACTION_WAYS[item["way"]]
        rows = [
            ("商户名称", merchant["name"], "", ""), ("统一社会信用代码", merchant["code"], "", ""),
            ("付款方", item["payer"], "", ""), ("收款方", item["recipient"], "", ""),
            ("交易金额", f'￥{item["amount"] / 100:.2f}', "交易日期", item["date"]),
            ("交易类型", "支出" if item["type"] == 0 else "收入", "交易方式", way),
            ("交易说明", item["description"], "", ""),
        ]
        info_rows = [ft.DataRow(cells=[ft.DataCell(ft.Text(a, weight=ft.FontWeight.W_600)), ft.DataCell(ft.Text(b)),
            ft.DataCell(ft.Text(c, weight=ft.FontWeight.W_600)), ft.DataCell(ft.Text(d))]) for a, b, c, d in rows]
        file_controls = []
        for index, file in enumerate(files, 1):
            path = Services.get_absolute_data_path(file["path"])
            controls: list[ft.Control] = [ft.Text(f'{index}. {path.name}', expand=True)]
            if Services.get_can_preview(path):
                controls.append(ft.IconButton(ft.Icons.VISIBILITY, tooltip="预览凭证",
                    on_click=lambda _e, p=path: self.set_preview_path(p)))
            controls.append(ft.IconButton(ft.Icons.FOLDER_OPEN, tooltip="打开文件所在位置",
                                          on_click=lambda _e, p=path: Services.set_open_folder(p)))
            file_controls.append(ft.Row(controls=controls))

        def set_modify(_event: ft.ControlEvent) -> None:
            self.page.pop_dialog()
            self.store.selected_merchant_id, self.selected_transaction_id = item["merchant_id"], item["id"]
            self.set_switch_view_by_title("收支记账")

        def set_delete(_event: ft.ControlEvent) -> None:
            self.page.pop_dialog()
            self.set_transaction_delete_dialog(item, False)

        self.page.show_dialog(ft.AlertDialog(modal=True, title="账目详情", scrollable=True,
            content=ft.Container(width=900, content=ft.Column(tight=True, controls=[
                ft.DataTable(columns=[ft.DataColumn(ft.Text("")) for _ in range(4)], rows=info_rows,
                             heading_row_height=0, column_spacing=20),
                ft.Divider(), ft.Text("相关凭证", weight=ft.FontWeight.W_600), *file_controls,
            ])), actions=[ft.Button("修改", on_click=set_modify), ft.Button("删除", color="#C62828", on_click=set_delete),
                         ft.TextButton("关闭", on_click=lambda _e: self.page.pop_dialog())]))

    def set_preview_path(self, path: Path) -> None:
        state = {"path": Services.get_relative_path(path), "temp_path": None}
        self.set_preview_file(state)

    def get_recent_quarters(self) -> list[tuple[str, int, int]]:
        current = date.today().year * 4 + (date.today().month - 1) // 3
        result = []
        for offset in range(6):
            value = current - offset
            year, quarter_index = divmod(value, 4)
            result.append((f"{year}年第{quarter_index + 1}季度", year, quarter_index + 1))
        return result

    def get_period_rows(self, merchant_id: int, start: date, end: date) -> list[dict[str, Any]]:
        return self.db.get_filtered_transactions(merchant_id, start.isoformat(), end.isoformat(), 2, 7, "date_asc")

    @staticmethod
    def get_totals(rows: list[dict[str, Any]]) -> tuple[int, int]:
        return (sum(x["amount"] for x in rows if x["type"] == 0),
                sum(x["amount"] for x in rows if x["type"] == 1))

    @staticmethod
    def get_daily_points(rows: list[dict[str, Any]], start: date, end: date, short: bool = False) -> list[tuple[str, int, int]]:
        values: dict[str, list[int]] = {}
        for row in rows:
            bucket = values.setdefault(row["date"], [0, 0])
            bucket[row["type"]] += row["amount"]
        points = []
        current = start
        while current <= end:
            value = values.get(current.isoformat(), [0, 0])
            points.append((current.strftime("%m-%d" if short else "%Y-%m-%d"), value[0], value[1]))
            current += timedelta(days=1)
        return points

    def get_stats_view(self) -> ft.Control:
        merchants = self.db.get_merchants()
        self.stats_merchant.options = [ft.DropdownOption(key=str(x["id"]), text=f'{x["name"]}<{x["code"]}>') for x in merchants]
        self.stats_merchant.value = str(self.store.selected_merchant_id) if self.store.selected_merchant_id else None
        self.stats_start.value, self.stats_end.value = self.accounts_start.value, self.accounts_end.value
        self.stats_quarter.options = [ft.DropdownOption(key=f"{year}-{quarter}", text=label)
                                      for label, year, quarter in self.get_recent_quarters()]
        if not self.stats_quarter.value:
            _, year, quarter = self.get_recent_quarters()[0]
            self.stats_quarter.value = f"{year}-{quarter}"
        self.set_query_stats(update=False)
        return ft.Column(expand=True, scroll=ft.ScrollMode.AUTO, spacing=16, controls=[
            ft.Row(controls=[self.stats_merchant, ft.Button("新增", icon=ft.Icons.ADD, on_click=self.set_jump_new_merchant)]),
            ft.Divider(),
            ft.Row(
                alignment=ft.MainAxisAlignment.SPACE_BETWEEN, spacing=8, controls=[
                    ft.Container(padding=ft.Padding(12, 6, 12, 6),
                                 border_radius=7, content=ft.Row(
                            spacing=8, controls=[
                                ft.Row(spacing=0, controls=[self.stats_start, self.get_stats_date_menu("start")]),
                                ft.Row(spacing=0, controls=[self.stats_end, self.get_stats_date_menu("end")]),
                                self.stats_year,self.stats_quarter
                            ])),
                    ft.Button("确认", icon=ft.Icons.CHECK, on_click=self.set_query_stats),
                ]), ft.Divider(), self.stats_cards, ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.SPACE_BETWEEN, controls=[
                ft.Text("收支时段统计", size=20, weight=ft.FontWeight.W_600),
                ft.Button("导出 Excel", icon=ft.Icons.DOWNLOAD, on_click=self.set_export_stats_excel)]),
            self.stats_donuts, ft.Divider(),
            ft.Row(alignment=ft.MainAxisAlignment.SPACE_BETWEEN, controls=[
                ft.Text("收支变化统计", size=20, weight=ft.FontWeight.W_600),
                ft.Button("导出图片", icon=ft.Icons.IMAGE, on_click=self.set_export_stats_images)]),
            self.stats_trends,
        ])

    def get_stats_date_menu(self, target: str) -> ft.PopupMenuButton:
        field = self.accounts_start if target == "start" else self.accounts_end
        stats_field = self.stats_start if target == "start" else self.stats_end
        original = field.value
        field.value = stats_field.value
        menu = self.get_date_menu(target)
        field.value = original
        for item in menu.items:
            old_click = item.on_click
            def wrap(handler, source=field, target_field=stats_field):
                def click(event):
                    old_start, old_end = self.accounts_start.value, self.accounts_end.value
                    self.accounts_start.value, self.accounts_end.value = self.stats_start.value, self.stats_end.value
                    handler(event)
                    target_field.value = source.value
                    self.accounts_start.value, self.accounts_end.value = old_start, old_end
                return click
            item.on_click = wrap(old_click)
        return menu

    def set_stats_merchant(self, event: ft.ControlEvent) -> None:
        self.store.selected_merchant_id = int(event.control.value) if event.control.value else None
        self.set_query_stats()

    def set_query_stats(self, _event: ft.ControlEvent | None = None, update: bool = True) -> None:
        self.accounts_start.value, self.accounts_end.value = self.stats_start.value, self.stats_end.value
        start, end = self.get_valid_filter_date(self.stats_start.value), self.get_valid_filter_date(self.stats_end.value)
        if not start:
            self.stats_start.value = self.accounts_start.value = ""
        if not end:
            self.stats_end.value = self.accounts_end.value = ""
        if start and end and start > end:
            start, end = end, start
            self.stats_start.value = self.accounts_start.value = start.isoformat()
            self.stats_end.value = self.accounts_end.value = end.isoformat()
        merchant_id = self.store.selected_merchant_id
        all_rows = self.db.get_transactions(merchant_id) if merchant_id else []
        actual_dates = [date.fromisoformat(x["date"]) for x in all_rows]
        period_start = start or (min(actual_dates) if actual_dates else date.today())
        period_end = end or (max(actual_dates) if actual_dates else date.today())
        period_rows = self.get_period_rows(merchant_id, period_start, period_end) if merchant_id else []
        year = int(self.stats_year.value)
        year_start, year_end = date(year, 1, 1), date(year, 12, 31)
        year_rows = self.get_period_rows(merchant_id, year_start, year_end) if merchant_id else []
        month_value = date.today().month if year == date.today().year else 1
        month_start = date(year, month_value, 1)
        month_end = date(year, month_value, monthrange(year, month_value)[1])
        month_rows = self.get_period_rows(merchant_id, month_start, month_end) if merchant_id else []
        quarter_year, quarter = map(int, self.stats_quarter.value.split("-"))
        quarter_start = date(quarter_year, (quarter - 1) * 3 + 1, 1)
        quarter_end_month = quarter * 3
        quarter_end = date(quarter_year, quarter_end_month, monthrange(quarter_year, quarter_end_month)[1])
        quarter_rows = self.get_period_rows(merchant_id, quarter_start, quarter_end) if merchant_id else []
        self.stats_data = {
            "period": (period_start, period_end, period_rows), "month": (month_start, month_end, month_rows),
            "quarter": (quarter_start, quarter_end, quarter_rows), "year": (year_start, year_end, year_rows),
        }
        self.set_render_stats()
        if update and self.page.controls:
            self.page.update()

    def get_week_points(self, rows: list[dict[str, Any]], start: date, end: date) -> list[tuple[str, int, int]]:
        weeks = (end - start).days // 7 + 1
        points = [[f"第{i + 1}周", 0, 0] for i in range(weeks)]
        for row in rows:
            index = (date.fromisoformat(row["date"]) - start).days // 7
            points[index][row["type"] + 1] += row["amount"]
        return [tuple(x) for x in points]

    def get_month_points(self, rows: list[dict[str, Any]]) -> list[tuple[str, int, int]]:
        points = [[f"{month}月", 0, 0] for month in range(1, 13)]
        for row in rows:
            month = date.fromisoformat(row["date"]).month
            points[month - 1][row["type"] + 1] += row["amount"]
        return [tuple(x) for x in points]

    def set_render_stats(self) -> None:
        period_rows = self.stats_data.get("period", (None, None, []))[2]
        expense, income = self.get_totals(period_rows)
        profit = income - expense
        profit_text = "0" if not profit else f'{"+" if profit > 0 else "-"}￥{abs(profit) / 100:.2f}'
        self.stats_cards.controls = [self.get_stat_card("已支出", f"-￥{expense / 100:.2f}", "#C62828"),
            self.get_stat_card("已收入", f"+￥{income / 100:.2f}", "#18864B"),
            self.get_stat_card("总利润", profit_text, "#18864B" if profit > 0 else "#C62828" if profit < 0 else "#202124")]
        if not self.store.selected_merchant_id or not any(value[2] for value in self.stats_data.values()):
            empty = self.empty_state(description="当前无账目数据，可前往创建", action_label="新增账目",
                                     action_icon=ft.Icons.ADD, action=self.set_jump_new_transaction)
            self.stats_donuts.controls, self.stats_trends.controls = [empty], []
            return
        labels = {
            "period": "时间段", "month": f'{self.stats_data["month"][0].month}月',
            "quarter": f'{(self.stats_data["quarter"][0].month - 1) // 3 + 1}季度',
            "year": f'{self.stats_year.value}年度',
        }
        self.stats_donuts.controls = []
        for key in ("period", "month", "quarter", "year"):
            e, i = self.get_totals(self.stats_data[key][2])
            controls: list[ft.Control] = [ft.Image(src=Services.get_donut_chart(e, i, labels[key]), fit=ft.BoxFit.CONTAIN)]
            self.stats_donuts.controls.append(ft.Container(col={"sm": 12, "md": 6, "xl": 3},
                alignment=ft.Alignment.CENTER, content=ft.Column(horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                                                                  controls=controls)))
        period_start, period_end, period_rows = self.stats_data["period"]
        month_start, month_end, month_rows = self.stats_data["month"]
        quarter_start, quarter_end, quarter_rows = self.stats_data["quarter"]
        _, _, year_rows = self.stats_data["year"]
        chart_specs = [
            (self.get_daily_points(period_rows, period_start, period_end), "时间段收支日变化"),
            (self.get_daily_points(month_rows, month_start, month_end, True), "月收支日变化"),
            (self.get_week_points(quarter_rows, quarter_start, quarter_end), "季度收支周变化"),
            (self.get_month_points(year_rows), "年度收支月变化"),
        ]
        self.stats_data["charts"] = chart_specs
        self.stats_trends.controls = []
        for chart_index, (points, title) in enumerate(chart_specs, 1):
            chart_bytes = Services.get_trend_chart(points, title)
            self.stats_trends.controls.append(ft.Container(
                alignment=ft.Alignment.CENTER,
                ink=True,
                on_click=lambda _e, data=chart_bytes, index=chart_index: self.set_chart_dialog(data, index),
                content=ft.Image(src=chart_bytes, fit=ft.BoxFit.CONTAIN),
            ))

    def set_chart_dialog(self, chart_bytes: bytes, chart_index: int) -> None:
        image = ft.Image(src=chart_bytes, fit=ft.BoxFit.CONTAIN)
        viewer = ft.InteractiveViewer(content=image, min_scale=0.2, max_scale=8, scale_factor=200,
                                      pan_enabled=True, scale_enabled=True)
        filename = f"qkbc_chart_{self.store.selected_merchant_id}_{chart_index}.webp"

        async def set_save(_event: ft.ControlEvent) -> None:
            output = BytesIO()
            with Image.open(BytesIO(chart_bytes)) as source:
                source.save(output, "WEBP", quality=92, method=6)
            saved = await ft.FilePicker().save_file(dialog_title="保存统计图", file_name=filename,
                                                    allowed_extensions=["webp"], src_bytes=output.getvalue())
            if saved:
                self.set_message("统计图已保存")

        self.page.show_dialog(ft.AlertDialog(
            modal=True,
            title="统计图预览",
            content=ft.Container(width=1050, height=620, bgcolor="#F7F9FA", alignment=ft.Alignment.CENTER,
                                 content=viewer),
            actions=[ft.Button("保存", icon=ft.Icons.SAVE, on_click=set_save),
                     ft.TextButton("关闭", on_click=lambda _e: self.page.pop_dialog())],
        ))

    async def set_export_stats_images(self, _event: ft.ControlEvent) -> None:
        if not self.store.selected_merchant_id or not self.stats_data.get("charts"):
            self.set_message("当前没有可导出的统计图", True)
            return
        folder = await ft.FilePicker().get_directory_path(dialog_title="选择统计图保存位置")
        if not folder:
            return
        for index, (points, title) in enumerate(self.stats_data["charts"], 1):
            png = Services.get_trend_chart(points, title)
            target = Path(folder) / f"qkbc_chart_{self.store.selected_merchant_id}_{index}.webp"
            with Image.open(BytesIO(png)) as chart:
                chart.save(target, "WEBP", quality=92, method=6)
        self.set_message("四张统计图已导出")

    @staticmethod
    def get_excel_period_values(rows: list[dict[str, Any]]) -> tuple[float | None, float | None, float | None]:
        if not rows:
            return None, None, None
        expense, income = QKBCASApp.get_totals(rows)
        return expense / 100, income / 100, (income - expense) / 100

    async def set_export_stats_excel(self, _event: ft.ControlEvent) -> None:
        if not self.store.selected_merchant_id:
            self.set_message("请选择商户", True)
            return
        merchant = self.db.get_merchant(self.store.selected_merchant_id)
        year = int(self.stats_year.value)
        period_start, period_end, period_rows = self.stats_data["period"]
        workbook = Workbook(); sheet = workbook.active; sheet.title = "账目统计"
        sheet.append(["商户", merchant["name"]])
        sheet.append(["统一社会信用代码", merchant["code"]])
        sheet.append(["开始日期", period_start.isoformat(), "结束日期", period_end.isoformat()])
        total_expense, total_income, total_profit = self.get_excel_period_values(period_rows)
        sheet.append(["总支出", total_expense, "总收入", total_income, "总利润", total_profit])
        year_rows = self.stats_data["year"][2]
        sheet.append([f"{year}年"])
        year_expense, year_income, year_profit = self.get_excel_period_values(year_rows)
        sheet.append(["年支出", year_expense, "年收入", year_income, "年利润", year_profit])
        for quarter in range(1, 5):
            quarter_start = date(year, (quarter - 1) * 3 + 1, 1)
            quarter_end_month = quarter * 3
            quarter_end = date(year, quarter_end_month, monthrange(year, quarter_end_month)[1])
            quarter_rows = self.get_period_rows(self.store.selected_merchant_id, quarter_start, quarter_end)
            sheet.append([f"{year}年第{quarter}季度"])
            e, i, p = self.get_excel_period_values(quarter_rows)
            sheet.append(["季度支出", e, "季度收入", i, "季度利润", p])
            for month in range((quarter - 1) * 3 + 1, quarter * 3 + 1):
                month_rows = [x for x in quarter_rows if date.fromisoformat(x["date"]).month == month]
                e, i, p = self.get_excel_period_values(month_rows)
                sheet.append([f"{month}月支出", e, f"{month}月收入", i, f"{month}月利润", p])
        for row in (1, 2, 3, 4, 5, 7, 12, 17, 22):
            for cell in sheet[row]:
                cell.font = Font(bold=True)
        for column in ("B", "D", "F"):
            for cell in sheet[column]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '¥#,##0.00;[Red]-¥#,##0.00'
        for column, width in zip("ABCDEF", [24, 26, 20, 18, 18, 18]):
            sheet.column_dimensions[column].width = width
        output = BytesIO(); workbook.save(output)
        filename = f"qkbc_statistics_{self.store.selected_merchant_id}_{year}.xlsx"
        saved = await ft.FilePicker().save_file(dialog_title="导出账目统计", file_name=filename,
            allowed_extensions=["xlsx"], src_bytes=output.getvalue())
        if saved:
            self.set_message("统计 Excel 已导出")

    def set_switch_view_by_title(self, title: str) -> None:
        index = next(i for i, item in enumerate(self.NAV_ITEMS) if item[0] == title)
        self.set_active_view(index)

    def set_jump_new_merchant(self, _event: ft.ControlEvent) -> None:
        self.set_switch_view_by_title("商户管理")
        self.set_new_merchant()

    def set_jump_new_transaction(self, _event: ft.ControlEvent) -> None:
        self.set_switch_view_by_title("收支记账")
        self.set_new_transaction(clear_merchant=False)

    def set_transaction_delete_dialog(self, item: dict[str, Any], second: bool) -> None:
        merchant = self.db.get_merchant(item["merchant_id"])

        def set_confirm(_event: ft.ControlEvent) -> None:
            self.page.pop_dialog()
            if not second:
                self.set_transaction_delete_dialog(item, True)
                return
            files = self.db.get_transaction_files(item["id"])
            try:
                self.db.set_delete_transaction(item["id"])
                for file in files:
                    path = Services.get_absolute_data_path(file["path"])
                    Services.set_remove_file(path)
                    Services.set_remove_empty_parents(path.parent)
                self.set_query_accounts()
                self.set_message("账目记录已删除")
            except Exception as error:
                self.set_message(f"删除失败：{error}", True)

        self.page.show_dialog(ft.AlertDialog(modal=True, title="再次确认删除" if second else "确认删除",
            content=ft.Column(tight=True, controls=[ft.Text(f'商户名称：{merchant["name"]}'),
                ft.Text(f'交易日期：{item["date"]}'), ft.Text(f'交易金额：￥{item["amount"] / 100:.2f}')]),
            actions=[ft.TextButton("取消", on_click=lambda _e: self.page.pop_dialog()),
                     ft.Button("确认", bgcolor="#C62828", color="#FFFFFF", on_click=set_confirm)]))

    async def set_export_accounts(self, _event: ft.ControlEvent) -> None:
        if not self.store.selected_merchant_id or not self.accounts_rows:
            self.set_message("当前没有可导出的账目记录", True)
            return
        workbook, sheet = Workbook(), None
        sheet = workbook.active
        sheet.title = "账目一览表"
        headers = ["序号", "交易日期", "交易类型", "金额", "说明", "交易方式", "收/付款方"]
        sheet.append(headers)
        for index, item in enumerate(self.accounts_rows, 1):
            way = item["way_remark"] if item["way"] == 6 else self.TRANSACTION_WAYS[item["way"]]
            party = item["recipient"] if item["type"] == 0 else item["payer"]
            sheet.append([index, item["date"], "支出" if item["type"] == 0 else "收入",
                          item["amount"] / 100, item["description"], way, party])
        for cell in sheet[1]:
            cell.font, cell.alignment = Font(bold=True), Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            row[4].alignment = Alignment(wrap_text=True, vertical="top")
        widths = [8, 14, 12, 14, 36, 18, 24]
        for column, width in zip("ABCDEFG", widths):
            sheet.column_dimensions[column].width = width
        output = BytesIO()
        workbook.save(output)
        start = self.get_valid_filter_date(self.accounts_start.value)
        end = self.get_valid_filter_date(self.accounts_end.value)
        suffix = date.today().strftime("%y%m%d")
        if start and end:
            suffix = start.strftime("%y%m%d") if start == end else f'{start:%y%m%d}_{end:%y%m%d}'
        elif start:
            suffix = start.strftime("%y%m%d")
        elif end:
            suffix = end.strftime("%y%m%d")
        filename = f"qkbc_transactions_{self.store.selected_merchant_id}_{suffix}.xlsx"
        saved = await ft.FilePicker().save_file(dialog_title="导出账目一览表", file_name=filename,
                                                allowed_extensions=["xlsx"], src_bytes=output.getvalue())
        if saved:
            self.set_message("Excel 已导出")

    def get_content_view(self, index: int) -> ft.Control:
        if index == 0:
            return self.get_merchant_view()
        if self.NAV_ITEMS[index][0] == "收支记账":
            return self.get_transaction_view()
        if self.NAV_ITEMS[index][0] == "所有账目":
            return self.get_accounts_view()
        if self.NAV_ITEMS[index][0] == "账目统计":
            return self.get_stats_view()
        return self.get_placeholder_view(index)

    def get_content_panel(self) -> ft.Container:
        return ft.Container(expand=True, padding=ft.Padding(30, 28, 30, 28), bgcolor="#F7F9FA",
                            content=self.get_content_view(self.store.get_selected_nav_index()))

    def set_refresh_merchants(self) -> None:
        merchants = self.db.get_merchants()
        self.merchant_dropdown.options = [
            ft.DropdownOption(key=str(item["id"]), text=f'{item["name"]}<{item["code"]}>') for item in merchants
        ]
        self.merchant_dropdown.value = str(self.store.selected_merchant_id) if self.store.selected_merchant_id else None

    def set_refresh_transaction_merchants(self) -> None:
        merchants = self.db.get_merchants()
        self.transaction_merchant_dropdown.options = [
            ft.DropdownOption(key=str(x["id"]), text=f'{x["name"]}<{x["code"]}>') for x in merchants
        ]
        self.transaction_merchant_dropdown.value = (
            str(self.store.selected_merchant_id) if self.store.selected_merchant_id else None
        )

    def set_transaction_merchant(self, event: ft.ControlEvent) -> None:
        self.store.selected_merchant_id = int(event.control.value) if event.control.value else None
        self.selected_transaction_id = None
        self.set_refresh_transactions()
        self.set_new_transaction(clear_merchant=False)

    def set_refresh_transactions(self) -> None:
        rows = self.db.get_transactions(self.store.selected_merchant_id) if self.store.selected_merchant_id else []
        self.transaction_dropdown.options = [
            ft.DropdownOption(key=str(x["id"]), text=f'{x["date"]}  {"收入" if x["type"] else "支出"}  ¥{x["amount"] / 100:.2f}')
            for x in rows
        ]
        self.transaction_dropdown.value = str(self.selected_transaction_id) if self.selected_transaction_id else None

    def set_select_transaction(self, event: ft.ControlEvent) -> None:
        if event.control.value:
            self.selected_transaction_id = int(event.control.value)
            self.set_load_transaction()

    def set_load_transaction(self) -> None:
        item = self.db.get_transaction(self.selected_transaction_id) if self.selected_transaction_id else None
        if not item:
            return
        self.amount_field.value = f'{item["amount"] / 100:.2f}'
        self.transaction_type.value = str(item["type"])
        self.payer_field.value, self.recipient_field.value = item["payer"], item["recipient"]
        self.way_dropdown.value, self.way_remark_field.value = str(item["way"]), item["way_remark"]
        self.way_remark_field.disabled = item["way"] != 6
        self.date_field.value, self.description_field.value = item["date"], item["description"]
        self.transaction_file_cards.clear()
        self.transaction_files_row.controls.clear()
        for saved in self.db.get_transaction_files(item["id"]):
            self.set_add_transaction_file_card(saved)
        self.page.update()

    def set_new_transaction(self, _event: ft.ControlEvent | None = None, clear_merchant: bool = False) -> None:
        self.selected_transaction_id = None
        self.transaction_dropdown.value = None
        self.amount_field.value, self.transaction_type.value = "", "0"
        self.payer_field.value = self.recipient_field.value = ""
        self.way_dropdown.value, self.way_remark_field.value = "0", ""
        self.way_remark_field.disabled = True
        self.date_field.value, self.description_field.value = date.today().isoformat(), ""
        self.transaction_file_cards.clear()
        self.transaction_files_row.controls.clear()
        if clear_merchant:
            self.store.selected_merchant_id = None
            self.transaction_merchant_dropdown.value = None
        if self.page.controls:
            self.page.update()

    def set_way_changed(self, event: ft.ControlEvent) -> None:
        self.way_remark_field.disabled = event.control.value != "6"
        if self.way_remark_field.disabled:
            self.way_remark_field.value = ""
        self.page.update()

    def set_update_payee_popup(self, event: ft.ControlEvent | None, category: int) -> None:
        field = self.payer_field if category == 1 else self.recipient_field
        popup = self.payer_popup if category == 1 else self.recipient_popup
        query = field.value or ""

        def get_select(name: str):
            def set_select(_event: ft.ControlEvent) -> None:
                field.value = name
                self.page.update()
            return set_select

        popup.items = [ft.PopupMenuItem(content=x["name"], on_click=get_select(x["name"]))
                       for x in self.db.get_payee_suggestions(query, category)]
        if event and popup.page:
            popup.update()

    def set_select_merchant(self, event: ft.ControlEvent) -> None:
        if event.control.value:
            self.store.selected_merchant_id = int(event.control.value)
            self.set_load_merchant()

    def set_load_merchant(self) -> None:
        merchant = self.db.get_merchant(self.store.selected_merchant_id) if self.store.selected_merchant_id else None
        if not merchant:
            return
        for field, key in [(self.name_field, "name"), (self.code_field, "code"), (self.address_field, "address"),
                           (self.legal_name_field, "legal_name"), (self.legal_phone_field, "legal_phone")]:
            field.value = merchant[key]
        self.file_cards.clear()
        self.attachments_row.controls.clear()
        for item in self.db.get_files(merchant["id"]):
            self.set_add_file_card(item)
        self.page.update()

    def set_new_merchant(self, _event: ft.ControlEvent | None = None) -> None:
        self.store.selected_merchant_id = None
        self.merchant_dropdown.value = None
        for field in [self.name_field, self.code_field, self.address_field, self.legal_name_field, self.legal_phone_field]:
            field.value = ""
            field.error_text = None
        self.file_cards.clear()
        self.attachments_row.controls.clear()
        self.page.update()

    def set_add_file_card(self, saved: dict[str, Any] | None = None) -> None:
        state: dict[str, Any] = {
            "id": saved.get("id") if saved else None,
            "path": saved.get("path") if saved else None,
            "temp_path": None,
            "original_path": saved.get("path") if saved else None,
        }
        type_dropdown = ft.Dropdown(
            label="文件类型", value=str(saved["type"]) if saved else "3",
            options=[ft.DropdownOption(key=str(k), text=v) for k, v in self.FILE_TYPES.items()], width=200,
        )
        name_field = ft.TextField(label="文件名称", hint_text="请输入文件名称", value=saved["name"] if saved else "", width=200)
        preview = ft.Container(width=200, height=200, bgcolor="#EDF1F3", border_radius=8, clip_behavior=ft.ClipBehavior.ANTI_ALIAS)
        menu = ft.Container(
            left=0,
            right=0,
            bottom=0,
            visible=True,
            bgcolor="#D91F2D33",
            padding=ft.Padding(4, 3, 4, 3),
            border_radius=ft.BorderRadius(0, 0, 8, 8),
        )
        state.update({"type_control": type_dropdown, "name_control": name_field, "preview": preview, "menu": menu})
        preview.content = self.get_file_preview_control(state)
        menu.content = self.get_file_menu(state)

        box = ft.Container(
            width=200,
            height=200,
            content=ft.Stack(expand=True, controls=[preview, menu]),
        )
        card = ft.Container(col={"sm": 12, "md": 6, "lg": 4, "xl": 3}, content=ft.Column(tight=True, spacing=8,
                            controls=[box, type_dropdown, name_field]))
        state["card"] = card
        self.file_cards.append(state)
        self.attachments_row.controls.append(card)
        if self.page.controls:
            self.page.update()

    def get_current_file(self, state: dict[str, Any]) -> Path | None:
        if state["temp_path"]:
            return Path(state["temp_path"])
        if state["path"]:
            return Services.get_absolute_data_path(state["path"])
        return None

    def get_file_preview_control(self, state: dict[str, Any]) -> ft.Control:
        path = self.get_current_file(state)
        if path and path.is_file() and Services.get_is_image(path):
            return ft.Image(src=str(path.resolve()), width=200, height=200, fit=ft.BoxFit.COVER)
        return ft.Container(alignment=ft.Alignment.CENTER, content=ft.Image(src="file.svg", width=72, height=72))

    def get_file_menu(self, state: dict[str, Any]) -> ft.Control:
        path = self.get_current_file(state)

        async def set_pick_file_click(_event: ft.ControlEvent) -> None:
            await self.set_pick_file(state)

        buttons = [ft.IconButton(
            ft.Icons.FIND_REPLACE if path else ft.Icons.UPLOAD_FILE,
            tooltip="替换任意文件" if path else "上传任意文件",
            icon_color="#FFFFFF",
            icon_size=21,
            on_click=set_pick_file_click,
        )]
        if path and Services.get_can_preview(path):
            buttons.append(ft.IconButton(ft.Icons.VISIBILITY, tooltip="预览", icon_color="#FFFFFF", icon_size=21,
                                         on_click=lambda _e: self.set_preview_file(state)))
        buttons.extend([
            ft.IconButton(ft.Icons.DELETE, tooltip="删除", icon_color="#FFFFFF", icon_size=21,
                          on_click=lambda _e: self.set_remove_card(state)),
            ft.IconButton(ft.Icons.FOLDER_OPEN, tooltip="打开文件位置", icon_color="#FFFFFF", icon_size=21,
                          disabled=not bool(path),
                          on_click=lambda _e: Services.set_open_folder(self.get_current_file(state))),
        ])
        return ft.Row(
            spacing=0,
            alignment=ft.MainAxisAlignment.CENTER,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
            controls=buttons,
        )

    async def set_pick_file(self, state: dict[str, Any]) -> None:
        files = await ft.FilePicker().pick_files(dialog_title="选择附件", allow_multiple=False)
        if not files or not files[0].path:
            return
        try:
            if state["temp_path"]:
                Services.set_remove_file(state["temp_path"])
            state["temp_path"] = str(Services.get_temp_copy(files[0].path))
            if not state["name_control"].value:
                state["name_control"].value = Path(files[0].name).stem
            state["preview"].content = self.get_file_preview_control(state)
            state["menu"].content = self.get_file_menu(state)
            self.page.update()
        except Exception as error:
            self.set_message(str(error), True)

    def set_remove_card(self, state: dict[str, Any]) -> None:
        if state["temp_path"]:
            Services.set_remove_file(state["temp_path"])
        self.file_cards.remove(state)
        self.attachments_row.controls.remove(state["card"])
        self.page.update()

    def set_add_transaction_file_card(self, saved: dict[str, Any] | None = None) -> None:
        state: dict[str, Any] = {
            "id": saved.get("id") if saved else None,
            "path": saved.get("path") if saved else None,
            "temp_path": None,
            "original_name": saved.get("original_name", "") if saved else "",
            "size": saved.get("size", 0) if saved else 0,
        }
        type_control = ft.Dropdown(label="文件类型", value=str(saved["type"]) if saved else "10", width=200,
            options=[ft.DropdownOption(key=str(k), text=v) for k, v in self.TRANSACTION_FILE_TYPES.items()])
        name_control = ft.TextField(label="文件名称", hint_text="请输入文件名称",
                                    value=saved["name"] if saved else "", width=200)
        preview = ft.Container(width=200, height=200, bgcolor="#EDF1F3", border_radius=8,
                               clip_behavior=ft.ClipBehavior.ANTI_ALIAS)
        menu = ft.Container(left=0, right=0, bottom=0, bgcolor="#D91F2D33", padding=ft.Padding(4, 3, 4, 3),
                            border_radius=ft.BorderRadius(0, 0, 8, 8))
        state.update({"type_control": type_control, "name_control": name_control, "preview": preview, "menu": menu})
        preview.content = self.get_file_preview_control(state)
        menu.content = self.get_transaction_file_menu(state)
        box = ft.Container(width=200, height=200, content=ft.Stack(expand=True, controls=[preview, menu]))
        card = ft.Container(col={"sm": 12, "md": 6, "lg": 4, "xl": 3},
                            content=ft.Column(tight=True, spacing=8, controls=[box, type_control, name_control]))
        state["card"] = card
        self.transaction_file_cards.append(state)
        self.transaction_files_row.controls.append(card)
        if self.page.controls:
            self.page.update()

    def get_transaction_file_menu(self, state: dict[str, Any]) -> ft.Control:
        path = self.get_current_file(state)

        async def set_pick(_event: ft.ControlEvent) -> None:
            files = await ft.FilePicker().pick_files(dialog_title="选择交易凭证", allow_multiple=False)
            if not files or not files[0].path:
                return
            try:
                if state["temp_path"]:
                    Services.set_remove_file(state["temp_path"])
                source = Path(files[0].path)
                state["temp_path"] = str(Services.get_temp_copy(source))
                state["original_name"], state["size"] = files[0].name, source.stat().st_size
                if not state["name_control"].value:
                    state["name_control"].value = source.stem
                state["preview"].content = self.get_file_preview_control(state)
                state["menu"].content = self.get_transaction_file_menu(state)
                self.page.update()
            except Exception as error:
                self.set_message(str(error), True)

        buttons = [ft.IconButton(ft.Icons.FIND_REPLACE if path else ft.Icons.UPLOAD_FILE,
                    tooltip="替换任意文件" if path else "上传任意文件", icon_color="#FFFFFF", on_click=set_pick)]
        if path and Services.get_can_preview(path):
            buttons.append(ft.IconButton(ft.Icons.VISIBILITY, tooltip="预览", icon_color="#FFFFFF",
                                         on_click=lambda _e: self.set_preview_file(state)))
        buttons.extend([
            ft.IconButton(ft.Icons.DELETE, tooltip="删除", icon_color="#FFFFFF",
                          on_click=lambda _e: self.set_remove_transaction_file_card(state)),
            ft.IconButton(ft.Icons.FOLDER_OPEN, tooltip="打开文件位置", icon_color="#FFFFFF", disabled=not bool(path),
                          on_click=lambda _e: Services.set_open_folder(self.get_current_file(state))),
        ])
        return ft.Row(spacing=0, alignment=ft.MainAxisAlignment.CENTER, controls=buttons)

    def set_remove_transaction_file_card(self, state: dict[str, Any]) -> None:
        if state["temp_path"]:
            Services.set_remove_file(state["temp_path"])
        self.transaction_file_cards.remove(state)
        self.transaction_files_row.controls.remove(state["card"])
        self.page.update()

    def set_save_transaction(self, _event: ft.ControlEvent) -> None:
        if not self.store.selected_merchant_id:
            self.set_message("请选择记账商户", True)
            return
        try:
            amount = int(round(float((self.amount_field.value or "").strip()) * 100))
            if amount < 0:
                raise ValueError
        except ValueError:
            self.set_message("请输入正确的金额", True)
            return
        payer, recipient = (self.payer_field.value or "").strip(), (self.recipient_field.value or "").strip()
        if not payer or not recipient:
            self.set_message("请输入付款方和收款方", True)
            return
        try:
            transaction_date = date.fromisoformat((self.date_field.value or "").strip()).isoformat()
        except ValueError:
            self.set_message("交易日期格式应为 YYYY-MM-DD", True)
            return
        way = int(self.way_dropdown.value)
        remark = (self.way_remark_field.value or "").strip()
        if way == 6 and not remark:
            self.set_message("交易方式为其他时必须填写备注", True)
            return
        for state in self.transaction_file_cards:
            if not self.get_current_file(state) or not (state["name_control"].value or "").strip():
                self.set_message("每个凭证都必须选择文件并填写文件名称", True)
                return
        merchant = self.db.get_merchant(self.store.selected_merchant_id)
        old_files = self.db.get_transaction_files(self.selected_transaction_id) if self.selected_transaction_id else []
        created: list[Path] = []
        try:
            payer_id = self.db.set_payee_history(payer, 1)
            recipient_id = self.db.set_payee_history(recipient, 2)
            values = {"merchant_id": self.store.selected_merchant_id, "type": int(self.transaction_type.value),
                      "amount": amount, "description": self.description_field.value or "", "way": way,
                      "way_remark": remark, "payer": payer, "recipient": recipient, "payer_id": payer_id,
                      "recipient_id": recipient_id, "date": transaction_date}
            transaction_id = self.db.set_save_transaction(self.selected_transaction_id, values)
            records, kept = [], set()
            for state in self.transaction_file_cards:
                file_type = int(state["type_control"].value)
                name = state["name_control"].value.strip()
                if state["temp_path"]:
                    target = Services.set_transaction_file(state["temp_path"], merchant["code"], transaction_id,
                        file_type, name, self.TRANSACTION_FILE_TYPES[file_type])
                    created.append(target)
                    relative = Services.get_relative_path(target)
                else:
                    relative = state["path"]
                    kept.add(relative)
                records.append({"original_name": state["original_name"], "name": name, "size": state["size"],
                                "type": file_type, "path": relative})
            self.db.set_replace_transaction_files(transaction_id, records)
            for item in old_files:
                if item["path"] not in kept and item["path"] not in {x["path"] for x in records}:
                    old_path = Services.get_absolute_data_path(item["path"])
                    Services.set_remove_file(old_path)
                    Services.set_remove_empty_parents(old_path.parent)
            for state in self.transaction_file_cards:
                if state["temp_path"]:
                    Services.set_remove_file(state["temp_path"])
            self.selected_transaction_id = transaction_id
            self.set_refresh_transactions()
            self.set_load_transaction()
            self.set_message("账目记录已保存")
        except Exception as error:
            for path in created:
                Services.set_remove_file(path)
            self.set_message(f"保存失败：{error}", True)

    def set_preview_file(self, state: dict[str, Any]) -> None:
        path = self.get_current_file(state)
        if not path or not path.is_file():
            return
        if path.suffix.lower() == ".pdf":
            self.set_pdf_dialog(path)
        else:
            self.page.show_dialog(ft.AlertDialog(modal=True, title="图片预览",
                content=ft.Container(width=900, height=620, content=ft.InteractiveViewer(
                    min_scale=0.2, max_scale=8, trackpad_scroll_causes_scale=True,
                    content=ft.Image(src=str(path.resolve()), fit=ft.BoxFit.CONTAIN))),
                actions=[ft.TextButton("关闭", on_click=lambda _e: self.page.pop_dialog())]))

    def set_pdf_dialog(self, path: Path) -> None:
        state = {"page": 0}
        preview_width, preview_height = self.get_pdf_preview_size()
        image = ft.Image(
            src=b"",
            width=preview_width,
            height=preview_height - 52,
            fit=ft.BoxFit.CONTAIN,
        )
        viewer = ft.InteractiveViewer(
            content=image,
            width=preview_width,
            height=preview_height - 52,
            min_scale=0.2,
            max_scale=8,
            trackpad_scroll_causes_scale=True,
        )
        counter, actions = ft.Text(), ft.Row(alignment=ft.MainAxisAlignment.CENTER)

        def get_page_buttons(total: int) -> list[ft.Control]:
            async def set_first(_event: ft.ControlEvent) -> None:
                await set_page(0)

            async def set_previous(_event: ft.ControlEvent) -> None:
                await set_page(state["page"] - 1)

            async def set_next(_event: ft.ControlEvent) -> None:
                await set_page(state["page"] + 1)

            async def set_last(_event: ft.ControlEvent) -> None:
                await set_page(total - 1)

            buttons: list[ft.Control] = []
            if state["page"] > 0:
                buttons.extend([
                    ft.IconButton(ft.Icons.FIRST_PAGE, tooltip="第一页", on_click=set_first),
                    ft.IconButton(ft.Icons.NAVIGATE_BEFORE, tooltip="上一页", on_click=set_previous),
                ])
            buttons.append(counter)
            if state["page"] < total - 1:
                buttons.append(ft.IconButton(ft.Icons.NAVIGATE_NEXT, tooltip="下一页", on_click=set_next))
                if state["page"] == 0:
                    buttons.append(ft.IconButton(ft.Icons.LAST_PAGE, tooltip="最后一页", on_click=set_last))
            return buttons

        async def set_page(index: int, update_controls: bool = True) -> None:
            data, total = Services.get_pdf_page(path, index)
            state["page"] = max(0, min(index, total - 1))
            image.src = data
            counter.value = f'{state["page"] + 1}/{total}'
            actions.controls = get_page_buttons(total)
            if update_controls:
                image.update()
                actions.update()
                await viewer.reset()

        data, total = Services.get_pdf_page(path, 0)
        image.src = data
        counter.value = f"1/{total}"
        actions.controls = get_page_buttons(total)
        preview_box = ft.Container(
            width=preview_width,
            height=preview_height,
            alignment=ft.Alignment.CENTER,
            content=ft.Column(
                spacing=4,
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                controls=[viewer, actions],
            ),
        )
        self.pdf_preview_box = preview_box
        self.pdf_preview_image = image
        self.pdf_preview_viewer = viewer

        def set_close(_event: ft.ControlEvent) -> None:
            self.pdf_preview_box = None
            self.pdf_preview_image = None
            self.pdf_preview_viewer = None
            self.page.pop_dialog()

        self.page.show_dialog(ft.AlertDialog(modal=True, title="PDF 预览",
            content=preview_box,
            actions=[ft.TextButton("关闭", on_click=set_close)]))

    def get_pdf_preview_size(self) -> tuple[int, int]:
        page_width = self.page.width or self.page.window.width or self.WINDOW_WIDTH
        page_height = self.page.height or self.page.window.height or self.WINDOW_HEIGHT
        return max(520, min(1000, int(page_width - 180))), max(400, min(720, int(page_height - 180)))

    def get_form_values(self) -> dict[str, str] | None:
        fields = [(self.name_field, "请输入商户名称"), (self.code_field, "请输入统一社会信用代码"),
                  (self.address_field, "请输入地址"), (self.legal_name_field, "请输入法人姓名"),
                  (self.legal_phone_field, "请输入法人手机号")]
        valid = True
        for field, message in fields:
            field.error_text = None if (field.value or "").strip() else message
            valid = valid and not field.error_text
        code = (self.code_field.value or "").strip().upper()
        self.code_field.value = code
        if code and not re.fullmatch(r"[A-Z0-9]+", code):
            self.code_field.error_text, valid = "只能输入大写字母和数字", False
        self.page.update()
        if not valid:
            return None
        return {"name": self.name_field.value.strip(), "code": code, "address": self.address_field.value.strip(),
                "legal_name": self.legal_name_field.value.strip(), "legal_phone": self.legal_phone_field.value.strip()}

    def set_save(self, _event: ft.ControlEvent) -> None:
        values = self.get_form_values()
        if not values:
            return
        for state in self.file_cards:
            if not (state["name_control"].value or "").strip() or not self.get_current_file(state):
                self.set_message("每个附件都必须选择文件并填写文件名称", True)
                return
        old_merchant = self.db.get_merchant(self.store.selected_merchant_id) if self.store.selected_merchant_id else None
        created_paths: list[Path] = []
        try:
            merchant_id = self.db.set_save_merchant(self.store.selected_merchant_id, values)
            saved_files = []
            kept_old_paths = set()
            for state in self.file_cards:
                name = state["name_control"].value.strip()
                file_type = int(state["type_control"].value)
                if state["temp_path"]:
                    target = Services.set_attachment(state["temp_path"], values["code"], file_type, name)
                    created_paths.append(target)
                    relative = Services.get_relative_path(target)
                elif old_merchant and old_merchant["code"] != values["code"]:
                    source = Services.get_absolute_data_path(state["path"])
                    target = Services.set_attachment(source, values["code"], file_type, name)
                    created_paths.append(target)
                    relative = Services.get_relative_path(target)
                else:
                    relative = state["path"]
                    kept_old_paths.add(relative)
                saved_files.append({"name": name, "type": file_type, "path": relative})
            old_files = self.db.get_files(merchant_id)
            self.db.set_replace_files(merchant_id, saved_files)
            for item in old_files:
                if item["path"] not in kept_old_paths and item["path"] not in {x["path"] for x in saved_files}:
                    old_path = Services.get_absolute_data_path(item["path"])
                    Services.set_remove_file(old_path)
                    Services.set_remove_empty_parents(old_path.parent)
            for state in self.file_cards:
                if state["temp_path"]:
                    Services.set_remove_file(state["temp_path"])
            self.store.selected_merchant_id = merchant_id
            self.set_refresh_merchants(); self.set_load_merchant()
            self.set_message("商户信息已保存")
        except Exception as error:
            for path in created_paths:
                Services.set_remove_file(path)
            self.set_message(f"保存失败：{error}", True)

    def set_request_delete(self, _event: ft.ControlEvent) -> None:
        merchant = self.db.get_merchant(self.store.selected_merchant_id) if self.store.selected_merchant_id else None
        if not merchant:
            self.set_message("请先选择要删除的商户", True)
            return
        self.set_delete_dialog(merchant, False)

    def set_delete_dialog(self, merchant: dict[str, Any], second: bool) -> None:
        def set_confirm(_event: ft.ControlEvent) -> None:
            self.page.pop_dialog()
            if not second:
                self.set_delete_dialog(merchant, True)
                return
            try:
                files = self.db.get_files(merchant["id"])
                transaction_files = [
                    file
                    for transaction in self.db.get_transactions(merchant["id"])
                    for file in self.db.get_transaction_files(transaction["id"])
                ]
                self.db.set_delete_merchant(merchant["id"])
                for item in files + transaction_files:
                    path = Services.get_absolute_data_path(item["path"])
                    Services.set_remove_file(path)
                    Services.set_remove_empty_parents(path.parent)
                self.set_new_merchant(); self.set_refresh_merchants()
                self.set_message("商户及其相关数据已删除")
            except Exception as error:
                self.set_message(f"删除失败：{error}", True)

        self.page.show_dialog(ft.AlertDialog(modal=True, title="再次确认删除" if second else "确认删除",
            content=ft.Column(tight=True, controls=[ft.Text(f'商户名称：{merchant["name"]}'), ft.Text(f'统一社会信用代码：{merchant["code"]}')]),
            actions=[ft.TextButton("取消", on_click=lambda _e: self.page.pop_dialog()),
                     ft.Button("确认", bgcolor="#C62828", color="#FFFFFF", on_click=set_confirm)]))

    def set_message(self, message: str, error: bool = False) -> None:
        self.page.show_dialog(ft.AlertDialog(title="提示", content=ft.Text(message),
            icon=ft.Icon(ft.Icons.CLOSE if error else ft.Icons.SAVE, color="#C62828" if error else "#176B45"),
            actions=[ft.TextButton("确定", on_click=lambda _e: self.page.pop_dialog())]))

    def set_active_view(self, index: int) -> None:
        self.store.set_selected_nav_index(index)
        for i, item in enumerate(self.nav_controls):
            selected, text = i == index, item.content.controls[1]
            item.bgcolor = "#E8F2ED" if selected else ft.Colors.TRANSPARENT
            text.color, text.weight = ("#176B45", ft.FontWeight.W_600) if selected else ("#36434A", ft.FontWeight.W_400)
        self.content_panel.content = self.get_content_view(index)
        self.page.update()

    def set_responsive_layout(self, _event: ft.PageResizeEvent) -> None:
        self.nav_panel.width = self.get_nav_width()
        if self.pdf_preview_box and self.pdf_preview_image and self.pdf_preview_viewer:
            width, height = self.get_pdf_preview_size()
            canvas_height = height - 52
            self.pdf_preview_box.width, self.pdf_preview_box.height = width, height
            self.pdf_preview_image.width, self.pdf_preview_image.height = width, canvas_height
            self.pdf_preview_viewer.width, self.pdf_preview_viewer.height = width, canvas_height
        self.page.update()

    async def get_start(self) -> None:
        await self.get_setup_window()
        self.nav_panel, self.content_panel = self.get_nav_panel(), self.get_content_panel()
        self.page.on_resized = self.set_responsive_layout
        self.page.add(ft.Row(expand=True, spacing=0, vertical_alignment=ft.CrossAxisAlignment.STRETCH,
                             controls=[self.nav_panel, self.content_panel]))


async def main(page: ft.Page) -> None:
    await QKBCASApp(page).get_start()


if __name__ == "__main__":
    ft.run(main, assets_dir=str(Services.get_asset_path("").resolve()))
